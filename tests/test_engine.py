import gzip
import io
import json
import subprocess
import tarfile
import zipfile
from pathlib import Path
from threading import Event

import pytest
from openpyxl import Workbook

from lotfinder.archives import find_7zip
from lotfinder.models import ScanConfig
from lotfinder.rules import Matcher
from lotfinder.scanner import Scanner
from lotfinder.store import Store


def dictionary(path, words):
    book = Workbook()
    book.active.append(["关键词"])
    for word in words:
        book.active.append([word])
    book.save(path)
    return path


def scan(tmp_path, source, **kwargs):
    config = ScanConfig([source], tmp_path / "out", **kwargs)
    store = Store(tmp_path / "result.sqlite3")
    Scanner(config, store).run()
    return store


def hits(store):
    return {(r["rule"], r["keyword"]): r["total"] for r in store.db.execute("SELECT * FROM hits")}


def test_rules_exact_suffix_and_boundaries(tmp_path):
    matcher = Matcher(ScanConfig([], tmp_path))
    result = matcher.matches("A00001 AT0004.13 AT0004.14 AL0535 A00001.S46 "
                             "AC000061 AD016984 AF017249.05 APE10264 ANG12345 ANS99999 "
                             "AC0000619 XAPE10264 AT12345 A00001_more")
    assert result[("product_lot", "AT0004.13")] == 1
    assert result[("product_lot", "AT0004.14")] == 1
    assert result[("product_lot", "A00001")] == 2
    assert result[("npw_lot", "AC000061")] == 1
    assert result[("foup", "APE10264")] == 1
    assert sum(result.values()) == 12


def test_counts_not_capped_evidence_per_line(tmp_path):
    file = tmp_path / "中文.txt"
    file.write_text("\n".join(["APE10264 APE10264 中文内容"] * 11), encoding="utf-8")
    store = scan(tmp_path, file)
    assert hits(store)["foup", "APE10264"] == 22
    evidence = store.evidence(1, "foup", "APE10264")
    assert len(evidence) == 5
    assert [e["line"] for e in evidence] == [1, 2, 3, 4, 5]
    assert evidence[0]["occurrences"] == 2
    store.close()


def test_three_dictionaries_and_identity(tmp_path):
    files = {key: dictionary(tmp_path / (key + ".xlsx"), words) for key, words in {
        "product": [" E0005X ", "E0005X", None, "E0005T"],
        "eqid": ["AMHTSEM02"], "layer": ["ETCH_01", "A+B"]}.items()}
    matcher = Matcher(ScanConfig([], tmp_path, dictionaries=files))
    found = matcher.matches("E0005X E0005T AMHTSEM02 ETCH_01 A+B XE0005X E0005X_more")
    assert sum(found.values()) == 5
    assert matcher.dictionary_info["product"]["count"] == 2
    assert not matcher.matches("e0005x")
    insensitive = Matcher(ScanConfig([], tmp_path, dictionaries=files, case_sensitive=False))
    assert insensitive.matches("e0005x")["product", "E0005X"] == 1


def test_dictionary_error_not_silent(tmp_path):
    bad = tmp_path / "missing.xlsx"
    with pytest.raises(FileNotFoundError):
        Matcher(ScanConfig([], tmp_path, dictionaries={"product": bad}))
    empty = dictionary(tmp_path / "empty.xlsx", [])
    with pytest.raises(ValueError, match="为空"):
        Matcher(ScanConfig([], tmp_path, dictionaries={"product": empty}))


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "utf-16-le", "gb18030"])
def test_encodings(tmp_path, encoding):
    file = tmp_path / "文本.log"
    file.write_bytes("设备日志 APE10264\n结束".encode(encoding))
    store = scan(tmp_path, file)
    assert hits(store)["foup", "APE10264"] == 1
    assert "设备日志" in store.evidence(1, "foup", "APE10264")[0]["text"]
    store.close()


def test_csv_multiline_and_neighbors(tmp_path):
    file = tmp_path / "table.csv"
    file.write_text('设备,批次,状态,其他\nM01,"AT0004.13\nAT0004.13",正常,不应保存\n', encoding="utf-8")
    store = scan(tmp_path, file)
    evidence = store.evidence(1, "product_lot", "AT0004.13")[0]
    assert hits(store)["product_lot", "AT0004.13"] == 2
    assert evidence["row"] == 2
    assert evidence["column"] == "B"
    assert [c["header"] for c in evidence["cells"]] == ["设备", "批次", "状态"]
    assert evidence["cells"][1]["value"].replace("\r\n", "\n") == "AT0004.13\nAT0004.13"
    store.close()


def test_xlsx_hidden_sheet_and_per_workbook_limit(tmp_path):
    file = tmp_path / "test.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["前列", "批次", "后列"])
    for _ in range(4):
        sheet.append(["设备", "AT0004.13", "成功"])
    hidden = book.create_sheet("隐藏表")
    hidden.sheet_state = "hidden"
    hidden.append(["前列", "批次", "后列"])
    for _ in range(4):
        hidden.append(["设备", "AT0004.13", "成功"])
    book.save(file)
    store = scan(tmp_path, file)
    assert hits(store)["product_lot", "AT0004.13"] == 8
    evidence = store.evidence(1, "product_lot", "AT0004.13")
    assert len(evidence) == 5
    assert evidence[-1]["sheet"] == "隐藏表"
    assert evidence[-1]["row"] == 2
    store.close()


def test_missing_formula_cache_marked_incomplete(tmp_path):
    file = tmp_path / "formulas.xlsx"
    book = Workbook()
    book.active.append(["列"])
    book.active.append(['="AT0004.13"'])
    book.save(file)
    store = scan(tmp_path, file)
    record = list(store.files())[0]
    assert record["status"] == "部分完成"
    assert "无缓存" in record["error"]
    assert not hits(store)
    store.close()


def test_legacy_xls_and_offsets(tmp_path):
    import xlwt
    file = tmp_path / "legacy.xls"
    book = xlwt.Workbook()
    sheet = book.add_sheet("测试")
    sheet.write(0, 1, "Lot")
    sheet.write(3, 1, "AC000061")
    sheet.write(3, 2, "右侧")
    book.save(str(file))
    store = scan(tmp_path, file)
    evidence = store.evidence(1, "npw_lot", "AC000061")[0]
    assert evidence["column"] == "B"
    assert evidence["row"] == 4
    assert evidence["cells"][1]["header"] == "Lot"
    store.close()


def test_ods(tmp_path):
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableRow, TableCell
    from odf.text import P
    file = tmp_path / "test.ods"
    document = OpenDocumentSpreadsheet()
    table = Table(name="数据")
    for values in [["设备", "Lot", "状态"], ["机台", "AF017249.05", "正常"]]:
        row = TableRow()
        for value in values:
            cell = TableCell(valuetype="string")
            cell.addElement(P(text=value))
            row.addElement(cell)
        table.addElement(row)
    document.spreadsheet.addElement(table)
    document.save(str(file))
    store = scan(tmp_path, file)
    evidence = store.evidence(1, "npw_lot", "AF017249.05")[0]
    assert evidence["column"] == "B"
    assert evidence["cells"][2]["value"] == "正常"
    store.close()


def zip_bytes(name, data):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(name, data)
    return stream.getvalue()


def test_five_nested_mixed_archives_and_filepart(tmp_path):
    data = zip_bytes("日志.txt", "APE10264 中文\n".encode())
    tar = io.BytesIO()
    with tarfile.open(fileobj=tar, mode="w") as archive:
        entry = tarfile.TarInfo("inner.zip")
        entry.size = len(data)
        archive.addfile(entry, io.BytesIO(data))
    gzip_path = tmp_path / "bundle.tar.gz"
    gzip_path.write_bytes(gzip.compress(tar.getvalue()))
    seven = tmp_path / "level.7z"
    subprocess.run([find_7zip(), "a", str(seven), str(gzip_path)], check=True, capture_output=True)
    outer = tmp_path / "batch.filepart"
    outer.write_bytes(zip_bytes("level.7z", seven.read_bytes()))
    store = scan(tmp_path, outer, max_depth=5)
    assert hits(store)["foup", "APE10264"] == 1
    files = list(store.files())
    assert len(files) == 6
    assert files[-1]["depth"] == 5
    assert "日志.txt" in files[-1]["source"]
    assert store.summary()["incomplete"] == 0
    store.close()


def test_sixth_layer_is_reported_not_silently_missed(tmp_path):
    data = b"APE10264"
    name = "data.txt"
    for _ in range(6):
        data = zip_bytes(name, data)
        name = "nested.zip"
    file = tmp_path / name
    file.write_bytes(data)
    store = scan(tmp_path, file, max_depth=5)
    assert list(store.files())[-1]["status"] == "受限"
    assert not hits(store)
    store.close()


def test_archive_path_traversal_cannot_write_outside_temp(tmp_path):
    file = tmp_path / "traversal.zip"
    file.write_bytes(zip_bytes("../../escape.txt", b"APE10264"))
    store = scan(tmp_path, file)
    assert hits(store)["foup", "APE10264"] == 1
    assert not (tmp_path.parent / "escape.txt").exists()
    store.close()


def test_damaged_file_distinct_from_no_match(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    (source / "bad.zip").write_bytes(b"PK\x03\x04broken")
    (source / "clean.txt").write_text("no match", encoding="utf-8")
    store = scan(tmp_path, source)
    statuses = {r["name"]: r["status"] for r in store.files()}
    assert statuses == {"bad.zip": "失败", "clean.txt": "完成"}
    store.close()


def test_cancel_preserves_partial_counts(tmp_path):
    file = tmp_path / "large.txt"
    file.write_text("APE10264\n" * 20)
    cancel = Event()
    store = Store(tmp_path / "cancel.sqlite3")
    scanner = Scanner(ScanConfig([file], tmp_path / "out"), store, cancel)
    original = store.record
    def record(*args):
        original(*args)
        cancel.set()
    store.record = record
    state = scanner.run()
    assert "取消" in state
    assert hits(store)["foup", "APE10264"] == 1
    assert list(store.files())[0]["status"] == "已取消"
    store.close()


def test_separate_files_keep_five_each(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    for name in ("one.txt", "two.txt"):
        (source / name).write_text("AC000061\n" * 8)
    store = scan(tmp_path, source)
    assert store.summary()["total"] == 16
    for file in store.files():
        assert len(store.evidence(file["id"], "npw_lot", "AC000061")) == 5
    store.close()


def test_expansion_limit(tmp_path):
    file = tmp_path / "large.zip"
    file.write_bytes(zip_bytes("big.txt", b"A" * 20000))
    store = scan(tmp_path, file, max_expanded_bytes=1000)
    assert store.summary()["incomplete"] >= 1
    assert not hits(store)
    store.close()


def test_xlsx_renamed_filepart(tmp_path):
    file = tmp_path / "book.filepart"
    book = Workbook()
    book.active.append(["Lot"])
    book.active.append(["AT0004.13"])
    book.save(file)
    store = scan(tmp_path, file)
    assert hits(store)["product_lot", "AT0004.13"] == 1
    assert list(store.files())[0]["kind"] == "xlsx"
    store.close()


def test_twenty_files_and_output_exclusion(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    out = source / "reports"
    out.mkdir()
    (out / "old.txt").write_text("APE10264")
    for i in range(20):
        (source / f"{i:02d}.txt").write_text("APE10264\n" * 6)
    store = Store(tmp_path / "twenty.sqlite3")
    Scanner(ScanConfig([source], out), store).run()
    assert store.summary() == {"files": 20, "total": 120, "matched": 20, "incomplete": 0, "archives": 0}
    assert store.db.execute("SELECT COUNT(*) FROM evidence").fetchone()[0] == 100
    store.close()


def test_split_7z_once(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    original = tmp_path / "test.txt"
    original.write_bytes((b"APE10264\n" * 50))
    subprocess.run([find_7zip(), "a", "-v100b", "-mx0", str(source / "split.7z"), str(original)],
                   check=True, capture_output=True)
    assert len(list(source.glob("*.00*"))) > 1
    store = scan(tmp_path, source)
    assert store.summary()["total"] == 50
    assert store.summary()["incomplete"] == 0
    store.close()


def test_xls_offset_rows_and_columns_without_header(tmp_path):
    import xlwt
    file = tmp_path / "offset.xls"
    book = xlwt.Workbook()
    sheet = book.add_sheet("offset")
    sheet.write(5, 3, "APE10264")
    book.save(str(file))
    store = scan(tmp_path, file, header_row=0)
    evidence = store.evidence(1, "foup", "APE10264")[0]
    assert (evidence["column"], evidence["row"]) == ("D", 6)
    store.close()
