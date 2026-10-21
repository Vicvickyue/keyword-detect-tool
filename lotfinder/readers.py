import codecs
import csv
import io
import zipfile
from pathlib import Path

from charset_normalizer import from_bytes
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from python_calamine import CalamineWorkbook, SheetTypeEnum

from .models import LimitExceeded


TEXT_SUFFIXES = {".txt", ".log", ".xml", ".json", ".ini", ".cfg", ".yaml", ".yml"}
TABLE_SUFFIXES = {".csv", ".tsv", ".cxv"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".xlsb", ".ods"}
ARCHIVE_SUFFIXES = {".zip", ".7z", ".rar", ".tar", ".gz", ".tgz", ".bz2", ".xz", ".001"}


def detect(path: Path, name: str):
    with path.open("rb") as stream:
        head = stream.read(8192)
    suffix = Path(name).suffix.lower()
    if head.startswith(b"SQLite format 3\x00") or suffix in {".db3", ".sqlite", ".sqlite3", ".db"}:
        return "sqlite"
    if head.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
                if "xl/workbook.xml" in names:
                    return "xlsx"
                if "xl/workbook.bin" in names:
                    return "xlsb"
                if "mimetype" in names:
                    item = archive.getinfo("mimetype")
                    if item.file_size < 200 and b"opendocument.spreadsheet" in archive.read(item):
                        return "ods"
        except zipfile.BadZipFile:
            pass
        return "archive"
    if head.startswith((b"7z\xbc\xaf\x27\x1c", b"Rar!", b"\x1f\x8b", b"BZh", b"\xfd7zXZ\x00")):
        return "archive"
    if head[257:262] == b"ustar":
        return "archive"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"
    if suffix in ARCHIVE_SUFFIXES:
        return "archive"
    if suffix in EXCEL_SUFFIXES:
        return suffix[1:]
    if suffix in TABLE_SUFFIXES:
        return "csv"
    if head.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE, codecs.BOM_UTF8)):
        return "text"
    if b"\x00" in head:
        # BOM-less UTF-16 is common in equipment logs.
        if len(head) > 8 and (head[1::2].count(0) > len(head[1::2]) * .6
                              or head[::2].count(0) > len(head[::2]) * .6):
            return "text"
        return "unsupported"
    if suffix in TEXT_SUFFIXES or suffix == ".filepart" or not suffix:
        return "text"
    if head and sum(b < 32 and b not in (9, 10, 13) for b in head) / len(head) < .01:
        return "text"
    return "unsupported"


def sample_encoding(path):
    with path.open("rb") as f:
        sample = f.read(256 * 1024)
    if sample.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    if sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"
    if sample and b"\x00" in sample:
        if sample[1::2].count(0) > len(sample[1::2]) * .6:
            return "utf-16-le"
        if sample[::2].count(0) > len(sample[::2]) * .6:
            return "utf-16-be"
    for encoding in ("utf-8", "gb18030"):
        try:
            codecs.getincrementaldecoder(encoding)(errors="strict").decode(sample, final=False)
            return encoding
        except UnicodeDecodeError:
            pass
    guess = from_bytes(sample).best()
    if guess is None:
        raise ValueError("无法识别文本编码；文件可能是二进制或不完整。")
    return guess.encoding


def encoding_for(path, check=lambda: None):
    """Validate an encoding over the entire file before any matches are recorded."""
    check()
    preferred = sample_encoding(path)
    with path.open("rb") as stream:
        head = stream.read(4)
    # A BOM or UTF-16 layout is authoritative: do not reinterpret corrupt Unicode as GBK.
    fixed = head.startswith((codecs.BOM_UTF8, codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)) or preferred.startswith("utf-16")
    candidates = [preferred] if fixed else list(dict.fromkeys([preferred, "utf-8", "gb18030"]))
    for encoding in candidates:
        decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
        try:
            with path.open("rb") as stream:
                while True:
                    check()
                    chunk = stream.read(256 * 1024)
                    decoder.decode(chunk, final=not chunk)
                    if not chunk:
                        return encoding
        except UnicodeDecodeError:
            continue
    raise ValueError("文本编码校验失败：文件可能包含混合编码、损坏或不完整字符。"
                     "请从原软件重新导出，或确认原始编码后转换为 UTF-8 再检索；未忽略无法解码的内容。")


def text_records(path, check):
    encoding = encoding_for(path, check)
    with path.open("r", encoding=encoding, errors="strict", newline="") as stream:
        number = 0
        while True:
            check()
            # A pathological unbroken line must not exhaust memory or silently be cut off.
            line = stream.readline(8 * 1024**2 + 1)
            if not line:
                break
            if len(line) > 8 * 1024**2:
                raise LimitExceeded("单行超过 8 Mi 字符，文件未完整扫描。")
            number += 1
            line = line.rstrip("\r\n")
            yield line, {"type": "text", "line": number, "text": line, "encoding": encoding}


def value_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def table_records(rows, sheet, header_row, check):
    # rows are (physical row index, complete row values), including blank prefixes.
    headers = []
    buffered = []
    for number, values in rows:
        check()
        texts = [value_text(v) for v in values]
        if header_row and number <= header_row:
            buffered.append((number, texts))
            if number == header_row:
                headers = texts
            else:
                continue
            to_scan = buffered
            buffered = []
        else:
            to_scan = [(number, texts)]
        for row_number, row in to_scan:
            for col, text in enumerate(row):
                if not text:
                    continue
                cells = []
                for idx in range(max(0, col - 1), min(len(row), col + 2)):
                    letter = get_column_letter(idx + 1)
                    cells.append({"column": letter, "header": headers[idx] if idx < len(headers) else "",
                                  "value": row[idx], "matched": idx == col})
                yield text, {"type": "table", "sheet": sheet, "row": row_number,
                             "column": get_column_letter(col + 1), "cells": cells}
    # File shorter than the configured header row: still search its contents.
    if buffered:
        yield from table_records(buffered, sheet, 0, check)


def csv_records(path, name, header_row, check):
    csv.field_size_limit(8 * 1024**2)
    with path.open("r", encoding=encoding_for(path, check), errors="strict", newline="") as stream:
        sample = stream.read(65536)
        stream.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel_tab if name.lower().endswith(".tsv") else csv.excel
        reader = csv.reader(stream, dialect, strict=True)
        yield from table_records(enumerate(reader, 1), "CSV（记录号）", header_row, check)


def check_workbook_zip(path, max_expanded, max_members):
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > max_members or sum(e.file_size for e in entries) > max_expanded:
                raise LimitExceeded("表格内部展开体积或条目数超过限制。")


def excel_records(path, kind, header_row, check, warnings):
    if kind in {"xlsx", "xlsm"}:
        # File-like input permits true XLSX files with a .filepart extension.
        with path.open("rb") as stream:
            book = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in book.worksheets:
                    check()
                    sheet.reset_dimensions()
                    yield from table_records(enumerate(sheet.iter_rows(values_only=True), 1),
                                             sheet.title, header_row, check)
            finally:
                book.close()
        # Find absent formula caches without searching formula expressions or double counting.
        with path.open("rb") as a, path.open("rb") as b:
            formulas = load_workbook(a, read_only=True, data_only=False, keep_links=False)
            cached = load_workbook(b, read_only=True, data_only=True, keep_links=False)
            missing = 0
            try:
                for fs, cs in zip(formulas.worksheets, cached.worksheets):
                    fs.reset_dimensions()
                    cs.reset_dimensions()
                    for fr, cr in zip(fs.iter_rows(), cs.iter_rows()):
                        check()
                        for fc, cc in zip(fr, cr):
                            if fc.data_type == "f" and cc.value is None:
                                missing += 1
                if missing:
                    warnings.append(f"{missing} 个公式单元格无缓存结果，未检索这些结果；请在 Excel 重算并保存后重试。")
            finally:
                formulas.close()
                cached.close()
    else:
        with CalamineWorkbook.from_path(path) as book:
            for meta in book.sheets_metadata:
                if meta.typ != SheetTypeEnum.WorkSheet:
                    continue
                check()
                sheet = book.get_sheet_by_name(meta.name)
                # Calamine iter_rows retains empty leading rows but omits leading columns.
                _, start_col = sheet.start or (0, 0)
                rows = ((number, [""] * start_col + values)
                        for number, values in enumerate(sheet.iter_rows(), 1))
                yield from table_records(rows, meta.name, header_row, check)
        warnings.append("此格式按已保存的单元格值检索，不执行公式重算；公式缓存完整性未逐项验证。")
