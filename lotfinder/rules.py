import hashlib
import re
from collections import Counter
from pathlib import Path

import ahocorasick
from openpyxl import load_workbook

from .models import DICT_RULES, RULE_NAMES
from .readers import check_workbook_zip


PATTERNS = {
    "product_lot": r"(?:A0|AT|AL)[0-9]{4}(?:\.[A-Za-z0-9]+)*",
    "npw_lot": r"(?:AC|AD|AF)[0-9]{6}(?:\.[A-Za-z0-9]+)*",
    "foup": r"(?:APE|ANG|ANS)[0-9]{5}",
}


def load_dictionary(path: Path):
    # The agreed contract is first worksheet / first column / A2 onwards.
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"词库需使用 .xlsx 或 .xlsm 文件：{path.name}")
    with path.open("rb") as source:
        digest = hashlib.file_digest(source, "sha256").hexdigest()
    book = load_workbook(path, read_only=True, data_only=False)
    words = []
    try:
        sheet = book.worksheets[0]
        for row in sheet.iter_rows(min_row=2, max_col=1):
            cell = row[0]
            if cell.data_type == "f":
                raise ValueError(f"词库 {path.name} 的 {cell.coordinate} 是公式，请填写实际 ID。")
            if cell.value is not None:
                value = str(cell.value).strip()
                if value:
                    words.append(value)
        return list(dict.fromkeys(words)), {"path": str(path.resolve()), "sha256": digest,
                                          "sheet": sheet.title, "count": len(set(words))}
    finally:
        book.close()


def id_char(ch):
    return bool(ch) and ch.isascii() and (ch.isalnum() or ch in "_.")


class Matcher:
    def __init__(self, config):
        self.case_sensitive = config.case_sensitive
        self.whole_id = config.whole_id
        flags = 0 if config.case_sensitive else re.IGNORECASE | re.ASCII
        # Underscore is a delimiter for fixed Lot/Foup formats. Dot suffixes belong to Lot IDs.
        self.regexes = {
            rule: re.compile(r"(?<![A-Za-z0-9])" + pat + r"(?![A-Za-z0-9]|\.[A-Za-z0-9])", flags)
            for rule, pat in PATTERNS.items() if rule in config.enabled
        }
        self.dictionary_info = {}
        self.active = list(self.regexes)
        self.warnings = []
        self.automaton = ahocorasick.Automaton()
        lookup = {}
        for rule in DICT_RULES:
            if rule not in config.enabled:
                continue
            path = config.dictionaries.get(rule)
            if not path:
                self.warnings.append(f"{RULE_NAMES[rule]}：未配置词库，规则未启用。")
                continue
            if path.stat().st_size > config.max_file_bytes:
                raise ValueError(f"词库大小超过单文件限制：{path.name}")
            check_workbook_zip(path, config.max_expanded_bytes, config.max_files)
            words, info = load_dictionary(path)
            self.dictionary_info[rule] = info
            if not words:
                raise ValueError(f"{RULE_NAMES[rule]} 词库为空，请从 A2 开始填写关键词。")
            self.active.append(rule)
            for word in words:
                key = word if self.case_sensitive else word.upper()
                lookup.setdefault(key, set()).add(rule)
        for word, rules in lookup.items():
            self.automaton.add_word(word, (word, tuple(sorted(rules))))
        self.has_dictionary = bool(lookup)
        if self.has_dictionary:
            self.automaton.make_automaton()
        if not self.active:
            raise ValueError("没有可运行的检索规则，请配置词库或启用固定格式规则。")

    def matches(self, text):
        result = Counter()
        for rule, pattern in self.regexes.items():
            for match in pattern.finditer(text):
                value = match.group()
                result[(rule, value if self.case_sensitive else value.upper())] += 1
        if self.has_dictionary:
            # str.upper can change string length; boundaries must use the normalized string too.
            search = text if self.case_sensitive else text.upper()
            for end, (word, rules) in self.automaton.iter(search):
                start = end - len(word) + 1
                if self.whole_id:
                    before = search[start - 1] if start else ""
                    after = search[end + 1] if end + 1 < len(search) else ""
                    if id_char(before) or id_char(after):
                        continue
                for rule in rules:
                    result[(rule, word)] += 1
        return result
