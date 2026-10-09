"""Create synthetic examples and empty dictionary templates; no user data is copied."""
import argparse
import gzip
import io
import tarfile
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def workbook(path, header, values):
    book = Workbook()
    sheet = book.active
    sheet.title = "关键词"
    sheet.append([header])
    for value in values:
        sheet.append([value])
    sheet["A1"].font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
    sheet["A1"].fill = PatternFill("solid", fgColor="087F8C")
    sheet.column_dimensions["A"].width = 32
    sheet.freeze_panes = "A2"
    book.save(path)


def zipped(name, content):
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as stream:
        stream.writestr(name, content)
    return data.getvalue()


def create(root):
    templates = root / "词库模板"
    samples = root / "示例"
    libraries = samples / "示例词库"
    inputs = samples / "待检索文件"
    for folder in (templates, libraries, inputs):
        folder.mkdir(parents=True, exist_ok=True)
    for name, values in {"Product_ID": ["E0005X", "E0005T"], "EQID": ["AMHTSEM02"],
                         "Layer_Step": ["ETCH_01"]}.items():
        workbook(templates / (name + ".xlsx"), name, [])
        workbook(libraries / (name + ".xlsx"), name, values)
    lines = ["2026-10-09 示例设备日志（全部为合成数据）",
             "产品批次 A00001.S46，产品 E0005X，设备 AMHTSEM02，工序 ETCH_01",
             "NPW 批次 AC000061 / AD016984 / AF017249.05",
             "产品批次 AT0004.13 和 AT0004.14 分开统计"]
    lines += [f"第 {index} 次记录：载具 APE10264，状态正常" for index in range(1, 9)]
    (inputs / "设备日志.txt").write_text("\n".join(lines), encoding="utf-8")
    (inputs / "表格.csv").write_text("设备,批次,状态,产品\nAMHTSEM02,AT0004.13,正常,E0005X\n"
                                    "AMHTSEM02,AT0004.14,完成,E0005T\n", encoding="utf-8-sig")
    book = Workbook()
    sheet = book.active
    sheet.title = "生产记录"
    sheet.append(["设备", "批次", "状态", "载具"])
    sheet.append(["AMHTSEM02", "AL0535", "正常", "APE10264"])
    sheet.append(["AMHTSEM02", "A00001.S46", "完成", "ANG12345"])
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 22
    hidden = book.create_sheet("隐藏记录")
    hidden.sheet_state = "hidden"
    hidden.append(["设备", "批次", "状态"])
    hidden.append(["AMHTSEM02", "AT0004.13", "隐藏行也检索"])
    book.save(inputs / "生产记录.xlsx")
    data = b"Nested archive: AT0004.13 APE10264\n"
    name = "inside.log"
    for level in range(1, 6):
        data = zipped(name, data)
        name = f"level-{level}.zip"
    (inputs / "五层压缩.filepart").write_bytes(data)
    (inputs / "无匹配.txt").write_text("这是没有关键词的文件。", encoding="utf-8")
    (inputs / "示例.xml").write_text('<record lot="AC000061">E0005X</record>', encoding="utf-8")
    return inputs, libraries


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    create(parser.parse_args().output)
